import pptx

prs = pptx.Presentation(r"C:\GitHub\l-data-seT---ML\Comprehensive_Viva_Crop_Fertilizer_ML.pptx")
print(f"Total slides: {len(prs.slides)}")
assert len(prs.slides) == 30, f"Expected 30 slides, got {len(prs.slides)}"

for idx, slide in enumerate(prs.slides, 1):
    tables = []
    pics = []
    texts = []
    for s in slide.shapes:
        if s.has_text_frame:
            t = s.text_frame.text.strip()
            if t:
                texts.append(t[:60].replace('\n', ' '))
        if s.has_table:
            tables.append(f"{len(s.table.rows)}x{len(s.table.columns)}")
        if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            pics.append(f"Pic({s.width/914400:.1f}x{s.height/914400:.1f})")
    first_t = texts[0] if texts else "EMPTY"
    print(f"Slide {idx:02d}: Pics={len(pics)}, Tables={tables}, FirstText='{first_t}'")

print("\nALL 30 SLIDES VERIFIED SUCCESSFULLY!")
